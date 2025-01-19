import knime.scripting.io as knio
import re
from pptx import Presentation
import pandas as pd

global_id = 0 # tracking for what text goes where

input_file = knio.flow_variables["tempPPTXpath"]

def contains_meaningful_content(text):
    return bool(re.search(r'[a-zA-Z0-9]', text))

def extract_shape_text(shape):
    """Extract text from shape and add a unique ID."""
    global global_id
    text_content = []
    if hasattr(shape, "text_frame") and shape.text_frame:
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if run.text.strip() and contains_meaningful_content(run.text):
                    extracted_text = run.text # temp storage of original                   
                    run.text = f"{global_id}"# add id to the text in PowerPoint for traceability

                    text_content.append({"id": global_id, "text": extracted_text})
                    global_id += 1
    return text_content

def extract_table_text(table):
    """Extract text from table cells and add a unique ID."""
    global global_id
    text_content = []
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.text.strip() and contains_meaningful_content(run.text):
                        extracted_text = run.text # temp storage of original
                        run.text = f"{global_id}" # add id to the text in PowerPoint for traceability

                        text_content.append({"id": global_id, "text": extracted_text})
                        global_id += 1
    return text_content

def process_shapes_recursive(shapes):
    """Process shapes to extract text and assign IDs."""
    all_text = []
    for shape in shapes:
        if shape.has_text_frame:
            all_text.extend(extract_shape_text(shape))
        elif shape.has_table:
            all_text.extend(extract_table_text(shape.table))
        elif hasattr(shape, 'shapes'):
            all_text.extend(process_shapes_recursive(shape.shapes))
    return all_text

def process_presentation(input_file):
    print(f"Opening {input_file}")
    try:
        input_ppt = Presentation(input_file)
    except Exception as e:
        print(f"Error opening file {input_file}: {e}")
        return []

    all_text = []
    for slide in input_ppt.slides:
        all_text.extend(process_shapes_recursive(slide.shapes))

    # save the updated powerpoint to the temp file with text replaced by their unique id
    input_ppt.save(input_file)
    print("Saving modified temp file with ids")
    
    return all_text

def main():
    extracted_text = process_presentation(input_file)

    text_df = pd.DataFrame(extracted_text)

    knio.output_tables[0] = knio.Table.from_pandas(text_df)

main()
